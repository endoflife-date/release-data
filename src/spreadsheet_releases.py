import csv
import io
import logging
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from src.common import http
from src.common.endoflife import AutoConfig, ProductFrontmatter
from src.common.parsing import ValueExtractor
from src.common.releasedata import ProductData

"""Fetch release-level data from a CSV or XLSX spreadsheet."""


def _value_as_string(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)


def _rows(url: str, content: bytes) -> Iterator[dict[str, object]]:
    suffix = Path(url.split("?", maxsplit=1)[0]).suffix.lower()
    if suffix != ".xlsx":  # Assume CSV
        yield from csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        return

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [_value_as_string(value) for value in next(rows)]
        if any(header is None for header in headers):
            msg = "XLSX header row contains an empty column name"
            raise ValueError(msg)

        for values in rows:
            yield dict(zip(headers, values, strict=False))
    finally:
        workbook.close()


class _Extractor(ValueExtractor[dict[str, object]]):
    def extract_raw_value(self, entry: dict[str, object]) -> str | None:
        return _value_as_string(entry.get(self.selector))


def update(_product: ProductFrontmatter, config: AutoConfig) -> None:
    with ProductData(config.product) as product_data:
        response = http.fetch_url(config.url)
        extractors = {
            name: _Extractor(name, definition)
            for name, definition in config.data["fields"].items()
        }

        if "releaseCycle" not in extractors:
            message = "fields must define releaseCycle"
            raise ValueError(message)

        release_cycle_extractor = extractors.pop("releaseCycle")
        for entry in _rows(config.url, response.content):
            release_name = release_cycle_extractor.extract(entry)
            if release_name is None:
                continue

            release = product_data.get_release(str(release_name))
            for name, extractor in extractors.items():
                try:
                    value = extractor.extract(entry)
                    if value is None:
                        continue
                except ValueError as error:
                    logging.debug("skipping field %s for %s: %s", name, release, error)
                    continue

                release.set_field(name, value)
